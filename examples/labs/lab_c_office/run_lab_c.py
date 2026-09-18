# -*- coding: utf-8 -*-
"""
LAB C：Office 闭环演示
读取 input_data.csv 销售明细，一键产出：
  1) sales_report.xlsx  —— openpyxl 生成，含「明细 / 区域汇总 / 产品汇总」三张表
  2) sales_report.docx  —— python-docx 生成的 Word 分析报告
  3) talking_points.md   —— 汇报要点
运行日志写入 lab_c.log
"""
import csv
import logging
import os
from collections import defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from docx import Document
from docx.shared import Pt, Cm

BASE = os.path.dirname(os.path.abspath(__file__))
INPUT_CSV = os.path.join(BASE, "input_data.csv")
OUT_XLSX = os.path.join(BASE, "sales_report.xlsx")
OUT_DOCX = os.path.join(BASE, "sales_report.docx")
OUT_MD = os.path.join(BASE, "talking_points.md")
LOG_FILE = os.path.join(BASE, "lab_c.log")

# ---------- 日志：同时写控制台与文件 ----------
logger = logging.getLogger("lab_c")
logger.setLevel(logging.INFO)
logger.handlers.clear()
_fmt = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s", "%Y-%m-%d %H:%M:%S")
_fh = logging.FileHandler(LOG_FILE, mode="w", encoding="utf-8")
_fh.setFormatter(_fmt)
_sh = logging.StreamHandler()
_sh.setFormatter(_fmt)
logger.addHandler(_fh)
logger.addHandler(_sh)


def load_rows(path):
    """读取销售明细 CSV，返回 list[dict]，并派生销售额字段。"""
    rows = []
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for r in reader:
            qty = int(r["数量"])
            price = float(r["单价"])
            amount = round(qty * price, 2)
            rows.append({
                "日期": r["日期"],
                "区域": r["区域"],
                "产品": r["产品"],
                "销售员": r["销售员"],
                "数量": qty,
                "单价": price,
                "销售额": amount,
            })
    logger.info("读取明细 %d 行，来自 %s", len(rows), path)
    return rows


def summarize(rows):
    """按区域 / 按产品分别汇总。"""
    by_region = defaultdict(lambda: {"订单数": 0, "总数量": 0, "总销售额": 0.0})
    by_product = defaultdict(lambda: {"订单数": 0, "总数量": 0, "总销售额": 0.0})
    for r in rows:
        g = by_region[r["区域"]]
        g["订单数"] += 1
        g["总数量"] += r["数量"]
        g["总销售额"] += r["销售额"]
        p = by_product[r["产品"]]
        p["订单数"] += 1
        p["总数量"] += r["数量"]
        p["总销售额"] += r["销售额"]
    for g in list(by_region.values()) + list(by_product.values()):
        g["总销售额"] = round(g["总销售额"], 2)
    total_amount = round(sum(r["销售额"] for r in rows), 2)
    logger.info("汇总完成：总销售额 %.2f 元，覆盖 %d 个区域、%d 个产品",
                total_amount, len(by_region), len(by_product))
    return by_region, by_product, total_amount


def build_excel(rows, by_region, by_product, total_amount, path):
    """用 openpyxl 生成三表 Excel 报告。"""
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")
    center = Alignment(horizontal="center", vertical="center")

    def style_header(ws, ncol):
        for c in range(1, ncol + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center
        for c in range(1, ncol + 1):
            ws.column_dimensions[get_column_letter(c)].width = 16

    # Sheet1：明细
    ws1 = wb.active
    ws1.title = "明细"
    cols = ["日期", "区域", "产品", "销售员", "数量", "单价", "销售额"]
    ws1.append(cols)
    for r in rows:
        ws1.append([r[c] for c in cols])
    ws1.append([])
    ws1.append(["合计", "", "", "", sum(r["数量"] for r in rows), "", total_amount])
    style_header(ws1, len(cols))
    logger.info("Excel[明细] 写入 %d 行明细 + 合计行", len(rows))

    # Sheet2：区域汇总
    ws2 = wb.create_sheet("区域汇总")
    ws2.append(["区域", "订单数", "总数量", "总销售额(元)"])
    for region, g in sorted(by_region.items(), key=lambda x: -x[1]["总销售额"]):
        ws2.append([region, g["订单数"], g["总数量"], g["总销售额"]])
    ws2.append(["合计", sum(g["订单数"] for g in by_region.values()),
                sum(g["总数量"] for g in by_region.values()), total_amount])
    style_header(ws2, 4)
    logger.info("Excel[区域汇总] 写入 %d 个区域", len(by_region))

    # Sheet3：产品汇总
    ws3 = wb.create_sheet("产品汇总")
    ws3.append(["产品", "订单数", "总数量", "总销售额(元)"])
    for product, g in sorted(by_product.items(), key=lambda x: -x[1]["总销售额"]):
        ws3.append([product, g["订单数"], g["总数量"], g["总销售额"]])
    ws3.append(["合计", sum(g["订单数"] for g in by_product.values()),
                sum(g["总数量"] for g in by_product.values()), total_amount])
    style_header(ws3, 4)
    logger.info("Excel[产品汇总] 写入 %d 个产品", len(by_product))

    wb.save(path)
    logger.info("Excel 已保存：%s", path)


def build_docx(rows, by_region, by_product, total_amount, path):
    """用 python-docx 生成 Word 报告。"""
    doc = Document()
    # 默认字号
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(11)

    doc.add_heading("2026年7-8月销售数据分析报告", level=0)
    doc.add_paragraph(f"报告生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph(f"数据来源：input_data.csv，共 {len(rows)} 笔订单，"
                      f"总销售额 {total_amount:,.2f} 元。")

    doc.add_heading("一、区域销售排名", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "区域", "订单数", "总数量", "总销售额(元)"
    for region, g in sorted(by_region.items(), key=lambda x: -x[1]["总销售额"]):
        cells = table.add_row().cells
        cells[0].text = region
        cells[1].text = str(g["订单数"])
        cells[2].text = str(g["总数量"])
        cells[3].text = f"{g['总销售额']:,.2f}"

    doc.add_heading("二、产品销售排名", level=1)
    table2 = doc.add_table(rows=1, cols=4)
    table2.style = "Light Grid Accent 1"
    hdr = table2.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "产品", "订单数", "总数量", "总销售额(元)"
    for product, g in sorted(by_product.items(), key=lambda x: -x[1]["总销售额"]):
        cells = table2.add_row().cells
        cells[0].text = product
        cells[1].text = str(g["订单数"])
        cells[2].text = str(g["总数量"])
        cells[3].text = f"{g['总销售额']:,.2f}"

    top_region = max(by_region.items(), key=lambda x: x[1]["总销售额"])
    top_product = max(by_product.items(), key=lambda x: x[1]["总销售额"])
    doc.add_heading("三、结论", level=1)
    doc.add_paragraph(
        f"1. 区域维度：{top_region[0]} 区销售额最高，达 {top_region[1]['总销售额']:,.2f} 元，"
        f"占总销售额 {top_region[1]['总销售额']/total_amount*100:.1f}%。"
    )
    doc.add_paragraph(
        f"2. 产品维度：{top_product[0]} 贡献最大，销售额 {top_product[1]['总销售额']:,.2f} 元，"
        f"占 {top_product[1]['总销售额']/total_amount*100:.1f}%。"
    )
    doc.add_paragraph(
        f"3. 本报告所有数字均可在配套 sales_report.xlsx 中逐格复核，原始明细未被修改。"
    )
    doc.save(path)
    logger.info("Word 报告已保存：%s", path)


def build_md(by_region, by_product, total_amount, path):
    """生成汇报要点 Markdown。"""
    top_region = max(by_region.items(), key=lambda x: x[1]["总销售额"])
    top_product = max(by_product.items(), key=lambda x: x[1]["总销售额"])
    lines = [
        "# 销售汇报要点（3 分钟版）",
        "",
        f"- 周期：2026-07-01 至 2026-08-07，共 {sum(g['订单数'] for g in by_region.values())} 笔订单。",
        f"- 总销售额：**{total_amount:,.2f} 元**。",
        f"- 第一名区域：**{top_region[0]}**，销售额 {top_region[1]['总销售额']:,.2f} 元，"
        f"占比 {top_region[1]['总销售额']/total_amount*100:.1f}%。",
        f"- 第一名产品：**{top_product[0]}**，销售额 {top_product[1]['总销售额']:,.2f} 元，"
        f"占比 {top_product[1]['总销售额']/total_amount*100:.1f}%。",
        "",
        "## 下一步建议",
        "1. 对销售靠后的区域复盘渠道与定价。",
        "2. 对头部产品准备备货与促销方案。",
        "3. 下周期把客单价、复购率纳入同一分析模板。",
        "",
        "> 本要点由 run_lab_c.py 自动生成，数字与 sales_report.xlsx / sales_report.docx 完全一致。",
    ]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    logger.info("汇报要点已保存：%s", path)


def main():
    logger.info("=" * 60)
    logger.info("LAB C 开始运行")
    rows = load_rows(INPUT_CSV)
    by_region, by_product, total_amount = summarize(rows)
    build_excel(rows, by_region, by_product, total_amount, OUT_XLSX)
    build_docx(rows, by_region, by_product, total_amount, OUT_DOCX)
    build_md(by_region, by_product, total_amount, OUT_MD)
    logger.info("LAB C 全部产物生成完毕：")
    for p in (OUT_XLSX, OUT_DOCX, OUT_MD, LOG_FILE):
        logger.info("  - %s (%d bytes)", os.path.basename(p), os.path.getsize(p))
    logger.info("LAB C 运行成功")


if __name__ == "__main__":
    main()
